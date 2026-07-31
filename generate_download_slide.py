from PIL import Image, ImageDraw, ImageFont
import qrcode
from qrcode.image.styledpil import StyledPilImage
from qrcode.image.styles.moduledrawers import RoundedModuleDrawer
from pptx import Presentation
from pptx.util import Inches, Pt
import io, os

APK_URL = "https://mlharum.unitani.com/storage/apk/AiHarumanis.apk"
PPTX_IN  = "/home/zainal/innovation/MLharum/MLharum-Presentation.pptx"
PPTX_OUT = "/home/zainal/innovation/MLharum/MLharum-Presentation-v3.pptx"
SLIDE_W, SLIDE_H = 3560, 2000   # 17.8" x 10" @ 200 DPI

BG        = "#0d2b1a"
GREEN     = "#2e7d32"
ACCENT    = "#66bb6a"
WHITE     = "#ffffff"
SUBTEXT   = "#b2dfb4"

def load_font(path, size):
    try:
        return ImageFont.truetype(path, size)
    except OSError:
        return ImageFont.load_default()

BOLD  = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
REG   = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"

# ── QR code ──────────────────────────────────────────────────────────────────
qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_H,
                   box_size=14, border=3)
qr.add_data(APK_URL)
qr.make(fit=True)
qr_img = qr.make_image(image_factory=StyledPilImage,
                       module_drawer=RoundedModuleDrawer(),
                       fill_color=GREEN, back_color="white")
qr_img = qr_img.convert("RGB")
QR_SIZE = 780
qr_img = qr_img.resize((QR_SIZE, QR_SIZE), Image.LANCZOS)

# ── Canvas ────────────────────────────────────────────────────────────────────
canvas = Image.new("RGB", (SLIDE_W, SLIDE_H), BG)
draw   = ImageDraw.Draw(canvas)

# Subtle top accent bar
draw.rectangle([(0, 0), (SLIDE_W, 12)], fill=ACCENT)
draw.rectangle([(0, SLIDE_H - 12), (SLIDE_W, SLIDE_H)], fill=ACCENT)

# Left column — text
LEFT_CX = SLIDE_W // 4        # center of left half

f_tag   = load_font(BOLD, 44)
f_title = load_font(BOLD, 108)
f_name  = load_font(BOLD, 80)
f_sub   = load_font(REG,  52)
f_url   = load_font(REG,  36)
f_note  = load_font(REG,  40)

def centered_text(text, font, y, color, cx=None):
    cx = cx or LEFT_CX
    bbox = draw.textbbox((0, 0), text, font=font)
    w = bbox[2] - bbox[0]
    draw.text((cx - w // 2, y), text, fill=color, font=font)

# Tag pill
tag_text = "MUAT TURUN SEKARANG"
tb = draw.textbbox((0, 0), tag_text, font=f_tag)
tw, th = tb[2] - tb[0], tb[3] - tb[1]
pill_x = LEFT_CX - tw // 2 - 24
pill_y = 200
draw.rounded_rectangle([pill_x, pill_y, pill_x + tw + 48, pill_y + th + 20],
                        radius=20, fill=ACCENT)
draw.text((pill_x + 24, pill_y + 10), tag_text, fill=BG, font=f_tag)

centered_text("Ai-Harumanis", f_title, 310, WHITE)
centered_text("Aplikasi Petani Mango", f_name, 440, ACCENT)

# Divider
div_y = 570
draw.rectangle([(LEFT_CX - 180, div_y), (LEFT_CX + 180, div_y + 4)], fill=GREEN)

centered_text("Imbas kod QR untuk", f_sub, 610, SUBTEXT)
centered_text("muat turun aplikasi", f_sub, 675, SUBTEXT)

# Features
features = ["Scan buah Harumanis", "Anggaran tarikh tuai", "Analisis kematangan pulpa"]
fy = 800
for feat in features:
    draw.ellipse([(LEFT_CX - 220, fy + 10), (LEFT_CX - 195, fy + 35)], fill=ACCENT)
    fb = draw.textbbox((0, 0), feat, font=f_note)
    draw.text((LEFT_CX - 175, fy), feat, fill=WHITE, font=f_note)
    fy += 70

# UniMAP footer bottom-left
uni_font = load_font(REG, 36)
draw.text((60, SLIDE_H - 80), "UniMAP · Sistem AI Harumanis", fill=SUBTEXT, font=uni_font)

# ── Right column — QR ────────────────────────────────────────────────────────
RIGHT_CX = SLIDE_W * 3 // 4

# White card behind QR
CARD_PAD = 40
card_x1 = RIGHT_CX - QR_SIZE // 2 - CARD_PAD
card_y1 = (SLIDE_H - QR_SIZE) // 2 - CARD_PAD - 30
card_x2 = RIGHT_CX + QR_SIZE // 2 + CARD_PAD
card_y2 = card_y1 + QR_SIZE + CARD_PAD * 2 + 60 + 30
draw.rounded_rectangle([card_x1, card_y1, card_x2, card_y2], radius=32, fill="white")

qr_x = RIGHT_CX - QR_SIZE // 2
qr_y = card_y1 + CARD_PAD + 10
canvas.paste(qr_img, (qr_x, qr_y))

# "Percuma · Android" below QR inside card
note_text = "Percuma  ·  Android"
nb = draw.textbbox((0, 0), note_text, font=f_note)
nw = nb[2] - nb[0]
draw.text((RIGHT_CX - nw // 2, qr_y + QR_SIZE + 16), note_text, fill=GREEN, font=f_note)

# URL below card
url_text = "mlharum.unitani.com"
ub = draw.textbbox((0, 0), url_text, font=f_url)
uw = ub[2] - ub[0]
draw.text((RIGHT_CX - uw // 2, card_y2 + 24), url_text, fill=SUBTEXT, font=f_url)

# ── Save slide image to buffer ────────────────────────────────────────────────
buf = io.BytesIO()
canvas.save(buf, format="PNG", dpi=(200, 200))
buf.seek(0)

# ── Append to PPTX ───────────────────────────────────────────────────────────
prs = Presentation(PPTX_IN)
slide_layout = prs.slide_layouts[6]   # blank layout
new_slide = prs.slides.add_slide(slide_layout)

slide_w_emu = prs.slide_width
slide_h_emu = prs.slide_height
new_slide.shapes.add_picture(buf, 0, 0, slide_w_emu, slide_h_emu)

prs.save(PPTX_OUT)
print(f"Saved: {PPTX_OUT}  ({len(prs.slides)} slides total)")

# Also save slide preview
canvas.save("/home/zainal/innovation/MLharum/demo-output/download-slide-preview.png", dpi=(200,200))
print("Preview: demo-output/download-slide-preview.png")
